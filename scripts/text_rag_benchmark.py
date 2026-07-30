#!/usr/bin/env python3
"""Text-only end-to-end RAG benchmark on PDF and XLSX files.

This is a fallback benchmark when visual/ColQwen models cannot be downloaded.
It uses PyMuPDF + openpyxl for extraction, sentence-transformers for embeddings,
sqlite-vec for storage, and Ollama for answer generation.
"""

from __future__ import annotations

import argparse
import logging
import time
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
import openpyxl

from local_rag_stack.config import Settings
from local_rag_stack.embedding.text_embedder import SentenceTransformersEmbedder
from local_rag_stack.extraction.base import DocumentExtractor
from local_rag_stack.generation.generator import AnswerGenerator
from local_rag_stack.models import ExtractedDocument, TextChunk
from local_rag_stack.retrieval.hybrid import HybridRetriever
from local_rag_stack.storage.sqlite_vec_store import SQLiteVecStore

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s: %(message)s")
logger = logging.getLogger(__name__)


class PyMuPDFExtractor(DocumentExtractor):
    """Extract text from PDFs using PyMuPDF."""

    def extract(
        self,
        file_path: Path,
        *,
        document_id: str | None = None,
        document_name: str | None = None,
        metadata: dict[str, Any] | None = None,
    ) -> ExtractedDocument:
        file_path = Path(file_path)
        document_id = document_id or self.make_document_id(file_path)
        document_name = document_name or file_path.name
        doc = fitz.open(str(file_path))
        chunks: list[TextChunk] = []
        for page_no, page in enumerate(doc, start=1):
            text = page.get_text()
            for i, chunk_text in enumerate(self.split_text(text, chunk_size=512, chunk_overlap=64)):
                chunks.append(
                    TextChunk(
                        chunk_id=f"{document_id}_p{page_no}_c{i:04d}",
                        document_id=document_id,
                        document_name=document_name,
                        page_number=page_no,
                        text=chunk_text,
                    )
                )
        doc.close()
        return ExtractedDocument(
            document_id=document_id,
            document_name=document_name,
            source_path=file_path,
            chunks=chunks,
            metadata=metadata or {},
        )


class XLSXTextExtractor(DocumentExtractor):
    """Extract text from XLSX workbooks using openpyxl."""

    def extract(
        self,
        file_path: Path,
        *,
        document_id: str | None = None,
        document_name: str | None = None,
        metadata: dict[str, Any] | None = None,
    ) -> ExtractedDocument:
        file_path = Path(file_path)
        document_id = document_id or self.make_document_id(file_path)
        document_name = document_name or file_path.name
        wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
        chunks: list[TextChunk] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(" | ".join(str(cell) if cell is not None else "" for cell in row))
            sheet_text = f"\nSheet: {sheet_name}\n" + "\n".join(rows)
            for i, chunk_text in enumerate(self.split_text(sheet_text, chunk_size=512, chunk_overlap=64)):
                chunks.append(
                    TextChunk(
                        chunk_id=f"{document_id}_{sheet_name}_c{i:04d}",
                        document_id=document_id,
                        document_name=document_name,
                        page_number=None,
                        section_heading=sheet_name,
                        text=chunk_text,
                    )
                )
        wb.close()
        return ExtractedDocument(
            document_id=document_id,
            document_name=document_name,
            source_path=file_path,
            chunks=chunks,
            metadata=metadata or {},
        )


def extract_file(file_path: Path) -> ExtractedDocument:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return PyMuPDFExtractor().extract(file_path)
    if suffix in (".xlsx", ".xls"):
        return XLSXTextExtractor().extract(file_path)
    raise ValueError(f"Unsupported file type: {suffix}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Text-only RAG benchmark")
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--queries", nargs="+", default=None)
    parser.add_argument("--data-dir", default="./data/benchmark_text")
    parser.add_argument("--model", default="qwen3.6:latest", help="Ollama generation model")
    parser.add_argument("--top-k", type=int, default=5)
    args = parser.parse_args()

    settings = Settings(
        data_dir=Path(args.data_dir),
        sqlite_vec_path=Path(args.data_dir) / "benchmark_text.vec.sqlite",
        text_embed_model="BAAI/bge-small-en-v1.5",
        vision_embed_backend="none",
        rerank_backend="none",
        llm_model=args.model,
    )  # type: ignore[call-arg]
    settings.ensure_paths()

    logger.info("Initializing embedder and store")
    embedder = SentenceTransformersEmbedder(model=settings.text_embed_model)
    store = SQLiteVecStore(settings.sqlite_vec_path)
    retriever = HybridRetriever(store, text_search_k=10, rrf_k=60)
    generator = AnswerGenerator(model=settings.llm_model, host=settings.ollama_host, timeout=300.0)

    logger.info("==> Ingestion")
    total_chunks = 0
    total_embed_time = 0.0
    for file_path in args.files:
        logger.info("Extracting %s", file_path.name)
        start = time.perf_counter()
        doc = extract_file(file_path)
        extract_time = time.perf_counter() - start

        logger.info("Embedding %d chunks from %s", len(doc.chunks), file_path.name)
        start = time.perf_counter()
        items = embedder.embed_chunks(doc.chunks)
        embed_time = time.perf_counter() - start
        store.add(items)

        total_chunks += len(items)
        total_embed_time += embed_time
        logger.info(
            "Ingested %s: extract %.2fs, embed %.2fs, %d chunks",
            file_path.name, extract_time, embed_time, len(items)
        )

    queries = args.queries or [
        "What is the main topic of this document?",
        "What are the technical specifications?",
        "What financial figures are reported?",
    ]

    logger.info("==> Retrieval + Generation")
    results: list[dict[str, Any]] = []
    for query in queries:
        logger.info("Query: %s", query)
        start = time.perf_counter()
        q_vector = embedder.embed_query(query)
        embed_time = time.perf_counter() - start

        start = time.perf_counter()
        retrieved = retriever.retrieve(q_vector, top_k=args.top_k)
        retrieve_time = time.perf_counter() - start

        start = time.perf_counter()
        response = generator.generate(query, retrieved, include_images=False)
        gen_time = time.perf_counter() - start

        results.append({
            "query": query,
            "embed_seconds": embed_time,
            "retrieve_seconds": retrieve_time,
            "generate_seconds": gen_time,
            "answer": response.answer,
            "sources": [
                {
                    "document": r.document_name,
                    "page": r.page_number,
                    "snippet": (r.snippet or "")[:100],
                    "score": r.score,
                }
                for r in response.citations
            ],
        })

    logger.info("==> Summary")
    print(f"\nFiles ingested: {len(args.files)}")
    print(f"Total chunks: {total_chunks}")
    print(f"Total embed time: {total_embed_time:.2f}s")
    print(f"Avg embed time per chunk: {total_embed_time / max(1, total_chunks) * 1000:.1f}ms")
    print("\nQuery results:")
    for r in results:
        print(f"\nQ: {r['query']}")
        print(f"   embed={r['embed_seconds']:.2f}s retrieve={r['retrieve_seconds']:.2f}s generate={r['generate_seconds']:.2f}s")
        print(f"A: {r['answer'][:500]}{'...' if len(r['answer']) > 500 else ''}")
        print("Sources:")
        for s in r["sources"]:
            snippet = f" | {s['snippet']}" if s["snippet"] else ""
            page = f" page={s['page']}" if s["page"] else ""
            print(f"   {s['document']}{page}{snippet} (score: {s['score']:.3f})")

    store.close()


if __name__ == "__main__":
    main()
