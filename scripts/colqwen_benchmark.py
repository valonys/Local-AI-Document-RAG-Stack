#!/usr/bin/env python3
"""End-to-end ColQwen multi-vector ingest/retrieve benchmark.

This script renders pages from a PDF and sheets from an XLSX into images,
embeds them with ColQwen2, stores them in the MultiVectorStore, and runs
retrieval queries. It intentionally bypasses Docling so it can run with only
PyMuPDF, openpyxl, and matplotlib installed.
"""

from __future__ import annotations

import argparse
import io
import logging
import time
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from PIL import Image

from local_rag_stack.config import Settings
from local_rag_stack.embedding.visual_embedder import ColQwenEmbedder
from local_rag_stack.models import MultiVectorItem, PageImage
from local_rag_stack.storage.multivec_store import MultiVectorStore

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s: %(message)s")
logger = logging.getLogger(__name__)


def render_pdf_pages(pdf_path: Path, dpi: int = 150) -> list[PageImage]:
    """Render each page of a PDF to a PageImage."""
    doc = fitz.open(str(pdf_path))
    pages: list[PageImage] = []
    for page_no in range(len(doc)):
        page = doc.load_page(page_no)
        pix = page.get_pixmap(dpi=dpi)
        img_bytes = pix.tobytes("png")
        img = Image.open(io.BytesIO(img_bytes))
        doc_id = pdf_path.stem
        pages.append(
            PageImage(
                image_id=f"{doc_id}_page_{page_no + 1:04d}",
                document_id=doc_id,
                document_name=pdf_path.name,
                page_number=page_no + 1,
                image_bytes=img_bytes,
                width=img.width,
                height=img.height,
            )
        )
    doc.close()
    return pages


def render_xlsx_sheets(xlsx_path: Path, max_rows: int = 200, dpi: int = 150) -> list[PageImage]:
    """Render each sheet of an XLSX workbook to a PageImage using matplotlib."""
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True, read_only=True)
    pages: list[PageImage] = []
    doc_id = xlsx_path.stem
    zoom = dpi / 100.0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            rows.append([str(cell) if cell is not None else "" for cell in row])

        if not rows:
            continue

        # Simple matplotlib table render.
        n_rows = len(rows)
        n_cols = max(len(r) for r in rows)
        fig_height = max(4, n_rows * 0.25)
        fig_width = max(6, n_cols * 1.2)
        fig, ax = plt.subplots(figsize=(fig_width, fig_height))
        ax.axis("off")
        ax.axis("tight")
        table = ax.table(cellText=rows, loc="center", cellLoc="left")
        table.auto_set_font_size(False)
        table.set_fontsize(8)
        table.scale(1, 1.5)
        fig.canvas.draw()
        buf = io.BytesIO()
        plt.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
        plt.close(fig)
        buf.seek(0)
        img_bytes = buf.getvalue()
        img = Image.open(io.BytesIO(img_bytes))

        pages.append(
            PageImage(
                image_id=f"{doc_id}_{sheet_name}_page_0001",
                document_id=doc_id,
                document_name=xlsx_path.name,
                page_number=1,
                image_bytes=img_bytes,
                width=img.width,
                height=img.height,
                metadata={"sheet": sheet_name, "rows_rendered": n_rows},
            )
        )
    wb.close()
    return pages


def ingest_files(
    embedder: ColQwenEmbedder,
    store: MultiVectorStore,
    files: list[Path],
) -> dict[str, Any]:
    """Render, embed, and store pages from each file."""
    stats = {"files": 0, "pages": 0, "embed_seconds": 0.0}
    for file_path in files:
        logger.info("Rendering %s", file_path.name)
        start = time.perf_counter()
        if file_path.suffix.lower() == ".pdf":
            pages = render_pdf_pages(file_path)
        elif file_path.suffix.lower() in (".xlsx", ".xls"):
            pages = render_xlsx_sheets(file_path)
        else:
            logger.warning("Unsupported file type: %s", file_path.suffix)
            continue
        render_time = time.perf_counter() - start

        logger.info("Embedding %d pages from %s", len(pages), file_path.name)
        start = time.perf_counter()
        items = embedder.embed_pages(pages, batch_size=2)
        embed_time = time.perf_counter() - start

        store.add(items)
        stats["files"] += 1
        stats["pages"] += len(items)
        stats["embed_seconds"] += embed_time
        logger.info(
            "Ingested %s: render %.2fs, embed %.2fs, %d vectors/page",
            file_path.name,
            render_time,
            embed_time,
            len(items[0].embeddings) if items else 0,
        )
    return stats


def run_queries(
    embedder: ColQwenEmbedder,
    store: MultiVectorStore,
    queries: list[str],
    top_k: int = 5,
) -> list[dict[str, Any]]:
    """Embed queries and retrieve top-K pages."""
    results: list[dict[str, Any]] = []
    for query in queries:
        logger.info("Query: %s", query)
        start = time.perf_counter()
        q_vectors = embedder.embed_query(query)
        embed_time = time.perf_counter() - start

        start = time.perf_counter()
        hits = store.search(q_vectors, top_k=top_k)
        search_time = time.perf_counter() - start

        results.append(
            {
                "query": query,
                "embed_seconds": embed_time,
                "search_seconds": search_time,
                "hits": [
                    {
                        "id": h.id,
                        "document": h.document_name,
                        "page": h.page_number,
                        "sheet": h.metadata.get("sheet"),
                        "score": h.score,
                    }
                    for h in hits
                ],
            }
        )
    return results


def main() -> None:
    parser = argparse.ArgumentParser(description="ColQwen multi-vector benchmark")
    parser.add_argument("files", nargs="+", type=Path, help="PDF/XLSX files to ingest")
    parser.add_argument("--queries", nargs="+", default=None, help="Queries to run")
    parser.add_argument("--model", default="vidore/colqwen2-v1.0", help="ColQwen model")
    parser.add_argument("--device", default="mps", help="torch device")
    parser.add_argument("--data-dir", default="./data/benchmark", help="Data directory")
    parser.add_argument("--top-k", type=int, default=5, help="Top-K results")
    args = parser.parse_args()

    settings = Settings(
        data_dir=Path(args.data_dir),
        sqlite_vec_path=Path(args.data_dir) / "benchmark.vec.sqlite",
    )  # type: ignore[call-arg]
    settings.ensure_paths()

    store_path = settings.sqlite_vec_path.with_suffix(".multivec.jsonl")
    logger.info("Initializing ColQwen embedder (model=%s, device=%s)", args.model, args.device)
    start = time.perf_counter()
    embedder = ColQwenEmbedder(model=args.model, device=args.device)
    init_time = time.perf_counter() - start
    logger.info("ColQwen ready in %.2fs", init_time)

    store = MultiVectorStore(store_path)

    logger.info("==> Ingestion")
    ingest_stats = ingest_files(embedder, store, args.files)

    queries = args.queries or [
        "What is the main topic of this document?",
        "Show me the table with financial figures.",
    ]
    logger.info("==> Retrieval")
    query_results = run_queries(embedder, store, queries, top_k=args.top_k)

    logger.info("==> Summary")
    print(f"\nFiles ingested: {ingest_stats['files']}")
    print(f"Pages embedded: {ingest_stats['pages']}")
    print(f"Total embed time: {ingest_stats['embed_seconds']:.2f}s")
    print(f"Avg embed time per page: {ingest_stats['embed_seconds'] / max(1, ingest_stats['pages']):.2f}s")
    print("\nQuery results:")
    for r in query_results:
        print(f"\n  Q: {r['query']}")
        print(f"     embed={r['embed_seconds']:.2f}s search={r['search_seconds']:.2f}s")
        for h in r["hits"]:
            sheet = f" sheet={h['sheet']}" if h["sheet"] else ""
            print(f"     [{h['score']:.3f}] {h['document']} page={h['page']}{sheet} ({h['id']})")

    store.close()


if __name__ == "__main__":
    main()
